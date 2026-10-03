"""Formula-driven 24-month funding fallback; all commercial terms are inputs.

This creates an ordinary new workbook, never re-saves the protected template.
The audit spec is emitted from the design before the workbook is authored.
Native Excel must recalculate the candidate before outputs are used.
"""
import argparse, calendar, json, math, datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from openpyxl.workbook.properties import CalcProperties

SCENARIO_KEYS = ['fte_first6','fte_later18','loaded_monthly_cost','user_scale','token_cost_scale','other_cost_scale','vendor_buffer','contingency_rate','reserve_months']
SCENARIO_LABELS = ['前6月等效人力','后18月等效人力','综合人月成本（元）','平台云用户规模倍率','人均算力成本倍率','其他建设成本倍率','厂商合作现金预留（元）','常规支出风险预留比例','期末正常运营储备月数']
GLOBAL_KEYS = ['months','start_month','owner_cash','recovery_rate','collection_lag','vendor_fee','vendor_prepay','vendor_deposit','vendor_month','asset_life_months','rounding_step']
GLOBAL_LABELS = ['规划月数','起始月份（假设）','已确认自有现金（未录入为0）','算力成本回收比例（非套餐价）','收款延迟月数','非退还合作费（报价未录入）','可抵扣算力预付（报价未录入）','可退保证金（报价未录入）','厂商款支付月份','设备折旧月数','资金目标取整步长（元）']

def usage_costs(p):
    """Calibrate short tasks from measured tokens; professional use is a separate benchmark."""
    u=p['usage']; t=u['task_sample']; price=u['deepseek_pro_peak_cny']
    task=(t['input_tokens']/t['completed_tasks']*price['input_miss'] +
          t['output_tokens']/t['completed_tasks']*price['output'])/1000000
    return {
        'student':task*u['student']['sessions_per_month']*u['student']['tasks_per_session'],
        'teacher':task*u['teacher']['sessions_per_month']*u['teacher']['tasks_per_session'],
        'personal':u['professional_monthly_usd']*u['usd_cny'],
        'development':u['developer_active_day_usd']*u['development_active_days']*u['usd_cny'],
        'task':task,
    }

def projected(p, scenario, **changes):
    """Design-side cash arithmetic, independent of Excel cell formulas."""
    p = dict(p, **changes); c=p['base_costs']; cash=p['owner_cash']; wallet=0
    unit=usage_costs(p)
    months=[]; revenue=[]; minimum=cash; deposit=0; totals={k:0.0 for k in ['payroll','dev_compute','infrastructure','teaching','community','admin','specialist','customer_token','vendor_fee','equipment','cash_paid','cash_collected','non_token_ops']}
    for m in range(1,25):
        people=[round(n*scenario['user_scale']) if m>=p['usage']['scene_budget_month'] else 0 for n in p['cohort_quarters'][(m-1)//3]]
        token=sum(n*unit[k] for n,k in zip(people,['student','teacher','personal']))*scenario['token_cost_scale']
        fte=scenario['fte_first6'] if m<=6 else scenario['fte_later18']; scale=scenario['other_cost_scale']
        values={'payroll':fte*scenario['loaded_monthly_cost'],'dev_compute':fte*unit['development']*scale,
            'infrastructure':(c['infrastructure_first6'] if m<=6 else c['infrastructure_middle6'] if m<=12 else c['infrastructure_last12'])*scale,
            'teaching':(c['teaching_first3'] if m<=3 else c['teaching_middle9'] if m<=12 else c['teaching_last12'])*scale,
            'community':(c['community_first3'] if m<=3 else c['community_middle9'] if m<=12 else c['community_last12'])*scale,
            'admin':c['admin_monthly']*scale,'specialist':c['specialist_months'].get(str(m),0)*scale,
            'customer_token':token,'vendor_fee':p['vendor_fee'] if m==p['vendor_month'] else 0,
            'equipment':c['equipment_months'].get(str(m),0)*scale}
        non_token=sum(values[k] for k in ['payroll','dev_compute','infrastructure','teaching','community','admin','specialist'])
        topup=p['vendor_prepay'] if m==p['vendor_month'] else 0; guarantee=p['vendor_deposit'] if m==p['vendor_month'] else 0
        used=min(wallet+topup,token); wallet=wallet+topup-used; deposit+=guarantee
        revenue.append(token*p['recovery_rate']); receipt=revenue[m-1-p['collection_lag']] if m>p['collection_lag'] else 0
        paid=non_token+token-used+topup+values['vendor_fee']+guarantee+values['equipment']; cash+=receipt-paid; minimum=min(minimum,cash)
        for k,v in values.items(): totals[k]+=v
        totals['cash_paid']+=paid; totals['cash_collected']+=receipt; totals['non_token_ops']+=non_token
        months.append({'month':m,'people':people,'token_cost':token,'cash_before_funding':cash,'cash_paid':paid,'cash_collected':receipt,'wallet':wallet,'deposit':deposit,'recurring':non_token-values['specialist']})
    contingency=totals['non_token_ops']*scenario['contingency_rate']
    unused=max(0,scenario['vendor_buffer']-totals['vendor_fee']-deposit-wallet)
    reserve=scenario['reserve_months']*(months[-1]['recurring']+months[-1]['token_cost']*(1-p['recovery_rate']))
    deficit=max(0,-minimum); required=deficit+contingency+unused+reserve; proposal=math.ceil(required/p['rounding_step'])*p['rounding_step']
    return {'name':scenario['name'],'totals':totals,'months':months,'peak_cash_deficit':deficit,'contingency':contingency,'unused_vendor_buffer':unused,'minimum_cash_reserve':reserve,'required_additional_funding':required,'proposed_funding':proposal,'ending_cash_with_proposed_funding':cash+proposal,'rounding_headroom':proposal-required}

def build(p, path, spec_path):
    expected=[projected(p,s) for s in p['scenarios']]
    numeric=[]
    for i,e in enumerate(expected):
        cc=col(i+2)
        for r,k in [(6,('totals','customer_token')),(11,('totals','cash_paid')),(14,('peak_cash_deficit',)),(18,('required_additional_funding',)),(19,('proposed_funding',)),(20,('ending_cash_with_proposed_funding',))]:
            value=e
            for key in k:value=value[key]
            numeric.append({'left':f'Summary!{cc}{r}','expected':value,'tolerance':0.01})
    spec={'required_sheets':['Summary','Assumptions','Usage','Cohorts','Costs','Revenue','Income Statement','Balance Sheet','Cash Flow','Financing','Checks','Sources'], 'min_formula_count':4000,'allow_external_links':False,
          'input_cells':['Assumptions!B4','Assumptions!B5','Assumptions!B7','Assumptions!B8','Assumptions!B9','Assumptions!C19'],
          'formula_cells':['Summary!C18','Summary!C19','Cash Flow!Y17','Checks!B8'],
          'policy_inputs':[{'cell':'Assumptions!B11','label':'Equipment useful life'},{'cell':'Assumptions!B12','label':'Funding rounding step'},{'cell':'Usage!B5','label':'Dated USD CNY reference'},{'cell':'Usage!B18','label':'Conditional scenario budget start month'}],
          'status_cells':[{'cell':f'Checks!{col(i+2)}8','equals':'PASS'} for i in range(3)],'numeric_checks':numeric}
    spec_path.parent.mkdir(parents=True,exist_ok=True); spec_path.write_text(json.dumps(spec,ensure_ascii=False,indent=2),encoding='utf-8')
    wb=openpyxl.Workbook(); wb.remove(wb.active)
    names=spec['required_sheets']; sheets={n:wb.create_sheet(n) for n in names}
    wb.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True)
    wb.properties.creator='Blocky Studio'; wb.properties.title='Blocky Studio 24个月持续发展资金规划'
    def cell(sheet,r,c,v,input=False):
        x=sheets[sheet].cell(r,c,v); x.font=Font(name='Microsoft YaHei',size=10,color='0000FF' if input else '000000'); x.alignment=Alignment(vertical='center',wrap_text=True)
        if isinstance(v,(float,int)) or isinstance(v,str) and v.startswith('='):x.number_format='#,##0.00;[Red]-#,##0.00;0'
        return x
    def heading(sheet,r,text):
        cell(sheet,r,1,text); sheets[sheet].cell(r,1).font=Font(name='Microsoft YaHei',size=12,bold=True)
    def monthly_headers(sheet,r):
        cell(sheet,r,1,'月份')
        for m in range(1,25):x=cell(sheet,r,m+1,f'=EDATE(Assumptions!$B$3,{m-1})');x.number_format='yyyy-mm'
        cell(sheet,r,26,'24个月合计 / 期末')
    def monthly_row(sheet,r,label,formulas,total='sum'):
        cell(sheet,r,1,label)
        for m,f in enumerate(formulas,2):cell(sheet,r,m,f)
        cell(sheet,r,26,f'=SUM(B{r}:Y{r})' if total=='sum' else f'=Y{r}')
    def ref(n,r,c):return f"'{n}'!{col(c)}{r}"
    a=sheets['Assumptions']; heading('Assumptions',1,'人民币；蓝色为可调输入，全部未证实金额与规模是规划假设')
    for r,(key,label) in enumerate(zip(GLOBAL_KEYS,GLOBAL_LABELS),2):
        value=p[key]
        if key=='start_month':value=datetime.date.fromisoformat(value+'-01')
        x=cell('Assumptions',r,1,label);x=cell('Assumptions',r,2,value,True)
        if key=='start_month':x.number_format='yyyy-mm'
        if key=='recovery_rate':x.number_format='0%'
    heading('Assumptions',14,p['unknown_terms'])
    for i,s in enumerate(p['scenarios'],2):cell('Assumptions',16,i,s['name'])
    for r,(key,label) in enumerate(zip(SCENARIO_KEYS,SCENARIO_LABELS),17):
        cell('Assumptions',r,1,label)
        for i,s in enumerate(p['scenarios'],2):cell('Assumptions',r,i,s[key],True)
    for r,(label,v) in enumerate([('学生短任务参考月成本','=Usage!B19'),('教师短任务参考月成本','=Usage!B20'),('专业云创作者参考月成本','=Usage!B21'),('研发等效人力月度Agent预算','=Usage!B22')],29):cell('Assumptions',r,1,label);cell('Assumptions',r,2,v)
    u=p['usage']; t=u['task_sample']; price=u['deepseek_pro_peak_cny']
    heading('Usage',1,'Token成本：项目短任务样本与专业开发者公开统计分开使用')
    usage_inputs={3:('有回执的任务数',t['completed_tasks']),4:('实际响应数',t['responses']),5:('美元兑人民币参考中间价',u['usd_cny']),6:('样本输入Token总量',t['input_tokens']),7:('样本输出Token总量',t['output_tokens']),8:('Pro高峰未命中输入价（元/百万）',price['input_miss']),9:('Pro高峰输出价（元/百万）',price['output']),10:('学生每月活动次数（假设）',u['student']['sessions_per_month']),11:('学生每活动任务尝试数（假设）',u['student']['tasks_per_session']),12:('教师每月准备次数（假设）',u['teacher']['sessions_per_month']),13:('教师每次任务尝试数（假设）',u['teacher']['tasks_per_session']),14:('专业开发者典型月成本下限（美元）',u['professional_monthly_usd_low']),15:('专业云创作者月成本中点（美元）',u['professional_monthly_usd']),16:('专业开发者典型月成本上限（美元）',u['professional_monthly_usd_high']),17:('研发每活跃日成本参考（美元）',u['developer_active_day_usd']),18:('场景预算启用月序（条件假设）',u['scene_budget_month']),23:('研发每月活跃天数（假设）',u['development_active_days'])}
    for r,(label,value) in usage_inputs.items():cell('Usage',r,1,label);cell('Usage',r,2,value,True)
    for r,label,f in [(19,'学生人均月短任务参考成本','=B24*B10*B11'),(20,'教师人均月短任务参考成本','=B24*B12*B13'),(21,'专业云创作者人均月参考成本','=B15*B5'),(22,'研发等效资源人均月Agent预算','=B17*B23*B5'),(24,'单任务按直供Pro峰价重定价（元）','=(B6/B3*B8+B7/B3*B9)/1000000'),(25,'每任务平均输入Token','=B6/B3'),(26,'每任务平均输出Token','=B7/B3'),(27,'专业月成本下限折人民币','=B14*B5'),(28,'专业月成本上限折人民币','=B16*B5')]:cell('Usage',r,1,label);cell('Usage',r,2,f)
    usage_notes={30:'样本仅涵盖10个有可读回执的短任务、30次响应；未取得缓存拆分。其他任务可能另有费用，不用成功样本估算其零消耗。',31:'学生/教师频次为规划输入，短任务参考不代表真实课堂月成本或完整专业模板消耗；输入输出、上下文增长与重试须后续实测。',32:'专业档单列Claude Code企业API使用统计150–250美元/月，中点为200；不套用于所有个人用户或小学生，不代表Blocky实测或厂商报价。',33:'内部研发采用13美元/活跃日×22个活跃日的资源预算，活跃天数是工作负荷假设；公开典型月范围与活跃日均不能同时相加。',34:'DeepSeek价格来自当前中文API文档，缓存未知按全部未命中敏感性重定价；不能当成息壤实际支付账单。',35:'汇率参考2026-09-30中国货币网6.7351，非支付成交汇率；美元统计先换算人民币再进入本模型。',36:'前18个月为研发排期假设，场景预算第19月后且产品准备关口通过才启用；日期与人数均非上线承诺。关口延期时调整启用月与人数曲线。',37:'BYOK及本地用户不计平台采购人数。平台云算力专业档用户规模是条件容量，人头不等于下载、注册、订阅销量或既有客户。'}
    for r,label in usage_notes.items():cell('Usage',r,1,label)
    sheets['Usage'].column_dimensions['A'].width=62;sheets['Usage'].column_dimensions['B'].width=22
    for r in usage_notes:sheets['Usage'].merge_cells(start_row=r,start_column=1,end_row=r,end_column=4);sheets['Usage'].row_dimensions[r].height=42
    base_keys=['infrastructure_first6','infrastructure_middle6','infrastructure_last12','teaching_first3','teaching_middle9','teaching_last12','community_first3','community_middle9','community_last12','admin_monthly']
    cost_labels=['基础设施前6月/月','基础设施第7至12月/月','基础设施后12月/月','教学共创前3月/月','教学共创第4至12月/月','教学共创后12月/月','社区前3月/月','社区第4至12月/月','社区后12月/月','一般运营/月']
    for r,(key,label) in enumerate(zip(base_keys,cost_labels),34):cell('Assumptions',r,1,label);cell('Assumptions',r,2,p['base_costs'][key],True)
    for c,label in enumerate(['月序','学生云活跃人数','教师云活跃人数','个人云活跃人数'],1):cell('Assumptions',46,c,label)
    for m in range(1,25):
        cell('Assumptions',46+m,1,m)
        for c,v in enumerate(p['cohort_quarters'][(m-1)//3],2):cell('Assumptions',46+m,c,v,True)
    cell('Assumptions',46,6,'专项外部交付月份');cell('Assumptions',46,7,'专项支出')
    for r,(m,v) in enumerate(p['base_costs']['specialist_months'].items(),47):cell('Assumptions',r,6,int(m),True);cell('Assumptions',r,7,v,True)
    cell('Assumptions',53,6,'设备采购月份');cell('Assumptions',53,7,'设备现金')
    for r,(m,v) in enumerate(p['base_costs']['equipment_months'].items(),54):cell('Assumptions',r,6,int(m),True);cell('Assumptions',r,7,v,True)
    rate=DataValidation(type='decimal',operator='between',formula1=0,formula2=1);a.add_data_validation(rate);rate.add(a['B5'])
    lag=DataValidation(type='whole',operator='between',formula1=0,formula2=6);a.add_data_validation(lag);lag.add(a['B6'])
    records=[]
    for i,s in enumerate(p['scenarios']):
        ac=col(i+2); co=1+i*14; ct=1+i*25; rv=1+i*12; cf=1+i*33; ip=1+i*18; bs=1+i*18
        for name,start in [('Cohorts',co),('Costs',ct),('Revenue',rv),('Cash Flow',cf),('Income Statement',ip),('Balance Sheet',bs)]:heading(name,start,s['name']);monthly_headers(name,start+1)
        sc=lambda r:f'Assumptions!${ac}${r}'
        for off,source_c,label in [(2,2,'学生平台云算力条件容量'),(3,3,'教师平台云算力条件容量'),(4,4,'专业创作者平台云算力条件容量')]:monthly_row('Cohorts',co+off,label,[f'=IF({m}>=Usage!$B$18,ROUND(Assumptions!{col(source_c)}{46+m}*{sc(20)},0),0)' for m in range(1,25)])
        monthly_row('Cohorts',co+5,'合计云算力人头（非下载数）',[f'=SUM({col(m+1)}{co+2}:{col(m+1)}{co+4})' for m in range(1,25)],'last')
        for off,personoff,price_r,label in [(6,2,29,'学生月度算力采购'),(7,3,30,'教师月度算力采购'),(8,4,31,'个人月度算力采购')]:monthly_row('Cohorts',co+off,label,[f'={col(m+1)}{co+personoff}*Assumptions!$B${price_r}*{sc(21)}' for m in range(1,25)])
        monthly_row('Cohorts',co+9,'外部用户Token成本合计',[f'=SUM({col(m+1)}{co+6}:{col(m+1)}{co+8})' for m in range(1,25)])
        cost_rows=[('等效资源人力',lambda m,c:f'=IF({m}<=6,{sc(17)},{sc(18)})'),('综合人力现金',lambda m,c:f'={c}{ct+2}*{sc(19)}'),('内部研发Agent预算',lambda m,c:f'={c}{ct+2}*Assumptions!$B$32*{sc(22)}'),('云基础设施',lambda m,c:f'=IF({m}<=6,Assumptions!$B$34,IF({m}<=12,Assumptions!$B$35,Assumptions!$B$36))*{sc(22)}'),('教学共创与试用',lambda m,c:f'=IF({m}<=3,Assumptions!$B$37,IF({m}<=12,Assumptions!$B$38,Assumptions!$B$39))*{sc(22)}'),('社区活动与材料',lambda m,c:f'=IF({m}<=3,Assumptions!$B$40,IF({m}<=12,Assumptions!$B$41,Assumptions!$B$42))*{sc(22)}'),('一般运营与合规预算',lambda m,c:f'=Assumptions!$B$43*{sc(22)}'),('专项外部交付',lambda m,c:f'=SUMIF(Assumptions!$F$47:$F$50,{m},Assumptions!$G$47:$G$50)*{sc(22)}'),('条件用户算力成本',lambda m,c:f'={ref("Cohorts",co+9,m+1)}'),('实际非退还合作费（未录入）',lambda m,c:f'=IF({m}=Assumptions!$B$10,Assumptions!$B$7,0)'),('费用现金合计（不含设备）',lambda m,c:f'=SUM({c}{ct+3}:{c}{ct+11})'),('设备购置现金',lambda m,c:f'=SUMIF(Assumptions!$F$54:$F$55,{m},Assumptions!$G$54:$G$55)*{sc(22)}'),('当月设备折旧',lambda m,c:f'=SUMPRODUCT(({m}-Assumptions!$F$54:$F$55+1>0)*({m}-Assumptions!$F$54:$F$55+1<=Assumptions!$B$11)*Assumptions!$G$54:$G$55)*{sc(22)}/Assumptions!$B$11'),('损益费用合计',lambda m,c:f'={c}{ct+12}+{c}{ct+14}'),('常规固定与专项支出',lambda m,c:f'=SUM({c}{ct+3}:{c}{ct+9})'),('常规经常性支出',lambda m,c:f'=SUM({c}{ct+3}:{c}{ct+8})')]
        for off,(label,fn) in enumerate(cost_rows,2):monthly_row('Costs',ct+off,label,[fn(m,col(m+1)) for m in range(1,25)])
        monthly_row('Revenue',rv+2,'条件用户算力成本',[f'={ref("Cohorts",co+9,m+1)}' for m in range(1,25)])
        monthly_row('Revenue',rv+3,'规划成本回收额（不是正式售价）',[f'={col(m+1)}{rv+2}*Assumptions!$B$5' for m in range(1,25)])
        monthly_row('Revenue',rv+4,'实际收款规划',[f'=IF({m}>Assumptions!$B$6,INDEX($B${rv+3}:$Y${rv+3},1,{m}-Assumptions!$B$6),0)' for m in range(1,25)])
        monthly_row('Revenue',rv+5,'期末待收回金额',[f'=SUM($B{rv+3}:{col(m+1)}{rv+3})-SUM($B{rv+4}:{col(m+1)}{rv+4})' for m in range(1,25)],'last')
        cf_labels=['月初现金（未补资）','收回的算力成本现金','常规固定与专项现金','用户Token消耗（费用口径）','厂商额度充值（现金时点）','月初预付额度资产','抵扣已购额度','月末预付额度资产','用户Token实际支付现金','非退还合作费用','可退保证金现金','设备购置现金','现金支付总计','当月净现金流','月末现金（未补资）','假设足额补入规划资金','月初现金（假设补足）','月末现金（假设补足）','保证金资产','额度资产','待收回金额','设备净值','充值/消耗现金恒等检查']
        for off,label in enumerate(cf_labels,2):cell('Cash Flow',cf+off,1,label)
        for m in range(1,25):
            c=col(m+1);prev=col(m); cur=lambda o:f'{c}{cf+o}'
            f={2:'=Assumptions!$B$4' if m==1 else f'={prev}{cf+16}',3:f'={ref("Revenue",rv+4,m+1)}',4:f'={ref("Costs",ct+16,m+1)}',5:f'={ref("Costs",ct+10,m+1)}',6:f'=IF({m}=Assumptions!$B$10,Assumptions!$B$8,0)',7:'=0' if m==1 else f'={prev}{cf+9}',8:f'=MIN({cur(7)}+{cur(6)},{cur(5)})',9:f'={cur(7)}+{cur(6)}-{cur(8)}',10:f'={cur(5)}-{cur(8)}+{cur(6)}',11:f'={ref("Costs",ct+11,m+1)}',12:f'=IF({m}=Assumptions!$B$10,Assumptions!$B$9,0)',13:f'={ref("Costs",ct+13,m+1)}',14:f'={cur(4)}+{cur(10)}+{cur(11)}+{cur(12)}+{cur(13)}',15:f'={cur(3)}-{cur(14)}',16:f'={cur(2)}+{cur(15)}',17:f'=IF({m}=1,Summary!{ac}$19,0)',18:'=Assumptions!$B$4' if m==1 else f'={prev}{cf+19}',19:f'={cur(18)}+{cur(3)}+{cur(17)}-{cur(14)}',20:f'=SUM($B{cf+12}:{c}{cf+12})',21:f'={cur(9)}',22:f'={ref("Revenue",rv+5,m+1)}',23:f'=SUM({ref("Costs",ct+13,2)}:{ref("Costs",ct+13,m+1)})-SUM({ref("Costs",ct+14,2)}:{ref("Costs",ct+14,m+1)})',24:f'={cur(10)}-{cur(5)}-{cur(6)}+{cur(8)}'}
            for off,v in f.items():cell('Cash Flow',cf+off,m+1,v)
        for off in range(2,25):cell('Cash Flow',cf+off,26,f'=Y{cf+off}' if off in [2,7,9,16,18,19,20,21,22,23] else f'=SUM(B{cf+off}:Y{cf+off})')
        for off,label in [(2,'规划成本回收收入'),(3,'外部用户算力成本'),(4,'算力贡献'),(5,'常规与专项运营费用'),(6,'非退还合作费'),(7,'折旧'),(8,'损益（零收入覆盖预算）')]:
            formulas=[]
            for m in range(1,25):
                c=col(m+1); f={2:f'={ref("Revenue",rv+3,m+1)}',3:f'={ref("Costs",ct+10,m+1)}',4:f'={c}{ip+2}-{c}{ip+3}',5:f'={ref("Costs",ct+16,m+1)}',6:f'={ref("Costs",ct+11,m+1)}',7:f'={ref("Costs",ct+14,m+1)}',8:f'={c}{ip+4}-{c}{ip+5}-{c}{ip+6}-{c}{ip+7}'};formulas.append(f[off])
            monthly_row('Income Statement',ip+off,label,formulas)
        for off,label in [(2,'现金（规划资金足额补入）'),(3,'待收回金额'),(4,'厂商预付额度资产'),(5,'保证金资产'),(6,'设备净值'),(7,'资产合计'),(8,'计划投入资本（非融资承诺）'),(9,'累计规划损益'),(10,'资产负债勾稽')]:
            fs=[]
            for m in range(1,25):
                c=col(m+1); f={2:f'={ref("Cash Flow",cf+19,m+1)}',3:f'={ref("Cash Flow",cf+22,m+1)}',4:f'={ref("Cash Flow",cf+21,m+1)}',5:f'={ref("Cash Flow",cf+20,m+1)}',6:f'={ref("Cash Flow",cf+23,m+1)}',7:f'=SUM({c}{bs+2}:{c}{bs+6})',8:f'=Assumptions!$B$4+Summary!{ac}$19',9:f'=SUM({ref("Income Statement",ip+8,2)}:{ref("Income Statement",ip+8,m+1)})',10:f'={c}{bs+7}-{c}{bs+8}-{c}{bs+9}'};fs.append(f[off])
            monthly_row('Balance Sheet',bs+off,label,fs,'last')
        records.append((co,ct,rv,cf,ip,bs))
    heading('Summary',1,'Blocky Studio｜24个月持续发展资金规划｜人民币')
    cell('Summary',2,1,'建设目标：个人开源、教育组织后台、多端、Agent App与社区；未报价参数均是规划假设。')
    for i,s in enumerate(p['scenarios'],2):cell('Summary',4,i,s['name'])
    labels={5:'期末平台云算力活跃人头',6:'用户Token采购（24月）',7:'研发工具与内部算力',8:'等效人力综合现金',9:'其他常规与专项运营',10:'设备现金',11:'24个月现金支付总额',12:'规划收款（默认不抵减）',13:'确认自有现金（未录入）',14:'峰值现金缺口',15:'常规支出风险预留',16:'未使用厂商合作现金预留',17:'正常运营现金储备',18:'最低追加资金需求',19:'建议资金规划目标（取整）',20:'假设足额补资的期末现金',21:'期末单月用户Token成本',22:'峰值缺口月份'}
    for r,label in labels.items():cell('Summary',r,1,label)
    for i,(co,ct,rv,cf,ip,bs) in enumerate(records):
        c=col(i+2);ac=c;sumcost=lambda r:f'SUM(Costs!B{r}:Y{r})'
        fs={5:f'=Cohorts!Y{co+5}',6:f'={sumcost(ct+10)}',7:f'={sumcost(ct+4)}',8:f'={sumcost(ct+3)}',9:f'={sumcost(ct+16)}-{c}7-{c}8',10:f'={sumcost(ct+13)}',11:f'=SUM(\'Cash Flow\'!B{cf+14}:Y{cf+14})',12:f'=SUM(Revenue!B{rv+4}:Y{rv+4})',13:'=Assumptions!$B$4',14:f'=MAX(0,-MIN(\'Cash Flow\'!B{cf+16}:Y{cf+16}))',15:f'={sumcost(ct+16)}*Assumptions!{ac}$24',16:f'=MAX(0,Assumptions!{ac}$23-{sumcost(ct+11)}-\'Cash Flow\'!Y{cf+20}-\'Cash Flow\'!Y{cf+21})',17:f'=Assumptions!{ac}$25*(Costs!Y{ct+17}+Costs!Y{ct+10}*(1-Assumptions!$B$5))',18:f'=SUM({c}14:{c}17)',19:f'=ROUNDUP({c}18/Assumptions!$B$12,0)*Assumptions!$B$12',20:f'=\'Cash Flow\'!Y{cf+19}',21:f'=Costs!Y{ct+10}',22:f'=TEXT(INDEX(\'Cash Flow\'!B{cf+1}:Y{cf+1},1,MATCH(MIN(\'Cash Flow\'!B{cf+16}:Y{cf+16}),\'Cash Flow\'!B{cf+16}:Y{cf+16},0)),"yyyy-mm")'}
        for r,f in fs.items():cell('Summary',r,i+2,f)
    heading('Summary',24,'默认成本回收率0：只表示未确认收入不用于压低资金，不能解释为实际商业收入为零。')
    heading('Summary',26,'厂商费用0表示未录入；现金预留是风险预算，获得报价后更新费、预付和保证金。')
    cell('Summary',28,1,'精简与高用量同时改变人力、人数和成本，是资源情景；新增用量须有额度与回款支持。')
    heading('Financing',1,'资金来源与用途｜全部融资仅为规划演算，不代表到账、估值或贷款')
    finance_labels={4:'假设外部资金补入',5:'确认自有资金',6:'规划运营收款',7:'现金来源合计',9:'研发/教研运营等效人力',10:'内部研发Agent预算',11:'外部用户Token实际现金',12:'基础设施/教学/社区/运营/专项',13:'设备采购',14:'合作费及保证金现金',15:'24个月现金支付',17:'期末现金（含储备与预留）',18:'现金用途合计',19:'来源减用途',21:'风险预留（不提前确认为费用）',22:'合作现金预留余额',23:'正常运营储备',24:'资金取整余量',26:'最低追加需求'}
    for r,label in finance_labels.items():cell('Financing',r,1,label)
    for i,(co,ct,rv,cf,ip,bs) in enumerate(records):
        c=col(i+2);cell('Financing',3,i+2,p['scenarios'][i]['name'])
        fs={4:f'=Summary!{c}19',5:f'=Summary!{c}13',6:f'=Summary!{c}12',7:f'=SUM({c}4:{c}6)',9:f'=Summary!{c}8',10:f'=Summary!{c}7',11:f'=SUM(\'Cash Flow\'!B{cf+10}:Y{cf+10})',12:f'=Summary!{c}9',13:f'=Summary!{c}10',14:f'=SUM(\'Cash Flow\'!B{cf+11}:Y{cf+12})',15:f'=SUM({c}9:{c}14)',17:f'=Summary!{c}20',18:f'={c}15+{c}17',19:f'={c}7-{c}18',21:f'=Summary!{c}15',22:f'=Summary!{c}16',23:f'=Summary!{c}17',24:f'=Summary!{c}19-Summary!{c}18',26:f'=Summary!{c}18'}
        for r,f in fs.items():cell('Financing',r,i+2,f)
    heading('Checks',1,'原生重算后的勾稽｜任何FAIL都阻止交付')
    for i,record in enumerate(records):
        co,ct,rv,cf,ip,bs=record;c=col(i+2);cell('Checks',3,i+2,p['scenarios'][i]['name'])
        fs={4:f'=ABS(Financing!{c}19)',5:f'=ABS(Summary!{c}19-Summary!{c}18-Financing!{c}24)',6:f'=MAX(MAX(\'Cash Flow\'!B{cf+24}:Y{cf+24}),-MIN(\'Cash Flow\'!B{cf+24}:Y{cf+24}))',7:f'=MAX(MAX(\'Balance Sheet\'!B{bs+10}:Y{bs+10}),-MIN(\'Balance Sheet\'!B{bs+10}:Y{bs+10}))',8:f'=IF(MAX({c}4:{c}7)<0.01,"PASS","FAIL")'}
        for r,f in fs.items():cell('Checks',r,i+2,f)
    for r,label in [(4,'现金来源与用途差额'),(5,'资金需求与取整差额'),(6,'预付抵扣恒等最大误差'),(7,'逐月资产负债最大误差'),(8,'状态')]:cell('Checks',r,1,label)
    heading('Sources',1,'参数来源和假设登记｜访问日期2026-10-03｜人民币')
    source_rows=[['F01','规划要求','24个月；Token按人头；合作费用未知','24个月持续发展规划','高（范围）',''],['F02','规划判断','第19月后条件容量/期末1160人','产品准备关口后；不代表上线或增长承诺','低',''],['F05','资源预算','3→4等效人力，综合人月2万元','覆盖研发、教研与运营，非现有人员履历','低',''],['F06','任务与公开使用统计','短任务实测Token与专业150–250美元/月分别使用','频次和专业预算中点为假设；不是售价或真实账单','中（来源）/低（频次）',''],['F07','厂商条件','实际费/额度预付/保证金待报价','0是未录入；合作现金预留30万元为预算判断','低',''],['F08','现金政策','3个月正常运营储备+10%常规支出风险预留','风险预留不含用户Token、设备及未知合作费','低',''],['F09','收入与资金','默认成本回收0、自有现金0','不以未确认收款或自有投入压低资金；外部补资仅规划','低',''],['F10','设备与会计','36个月折旧；预付和保证金为资产','现金、费用与折旧分别核算；无融资税务承诺','低','']]
    for source in p['sources'][2:]:source_rows.append([source['id'],source['basis'],source['description'],'来源类别及边界见原说明；USD按2026-09-30参考汇率换算','高（机制）',source['url']])
    for r,row in enumerate([['ID','类型','参数/事实','依据和边界','置信度','公开来源']]+source_rows,3):
        for c,v in enumerate(row,1):cell('Sources',r,c,v)
    for name,sheet in sheets.items():
        sheet.freeze_panes='B3';sheet.sheet_view.showGridLines=True
        sheet.column_dimensions['A'].width=62 if name=='Usage' else 39
        for c in range(2,sheet.max_column+1):sheet.column_dimensions[col(c)].width=13
        sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth=1 if name in ['Summary','Assumptions','Usage','Financing','Checks','Sources'] else 2;sheet.page_setup.fitToHeight=0;sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.print_options.horizontalCentered=True;sheet.print_title_cols='A:A';sheet.print_area=sheet.dimensions
        sheet.page_margins.left=0.25;sheet.page_margins.right=0.25
        for row in sheet:
            for x in row:
                if x.row in [1,3,4,16] and name in ['Summary','Financing','Assumptions','Checks','Sources']:x.fill=PatternFill('solid',fgColor='E5EBEE')
        sheet.oddFooter.center.text='Blocky Studio | 规划预算 | &P / &N'
    for n in ['Summary','Financing','Checks']:
        for c in ['B','C','D']:sheets[n].column_dimensions[c].width=22
        sheets[n].row_dimensions[1].height=30
        sheets[n].merge_cells('A1:D1')
    sheets['Summary'].merge_cells('A2:D2');sheets['Summary'].merge_cells('A24:D24');sheets['Summary'].merge_cells('A26:D26');sheets['Summary'].merge_cells('A28:D28');sheets['Summary'].row_dimensions[28].height=36
    sheets['Assumptions'].merge_cells('A1:G1');sheets['Assumptions'].merge_cells('A14:G14');sheets['Assumptions'].row_dimensions[1].height=32
    sheets['Sources'].merge_cells('A1:F1');sheets['Sources'].row_dimensions[1].height=32
    for name,stride in [('Cohorts',14),('Costs',25),('Revenue',12),('Income Statement',18),('Balance Sheet',18),('Cash Flow',33)]:
        for i in [1,2]:sheets[name].row_breaks.append(Break(id=i*stride))
        sheets[name].col_breaks.append(Break(id=13))
    for n in ['Assumptions','Summary','Financing','Checks']:
        for r in range(2,sheets[n].max_row+1):
            if sheets[n].row_dimensions[r].height is None:sheets[n].row_dimensions[r].height=19 if n=='Assumptions' else 24
    sheets['Summary'].column_dimensions['A'].width=62;sheets['Summary'].row_dimensions[2].height=50;sheets['Summary'].row_dimensions[24].height=45;sheets['Summary'].row_dimensions[26].height=45
    sheets['Sources'].column_dimensions['A'].width=10;sheets['Sources'].column_dimensions['B'].width=20;sheets['Sources'].column_dimensions['C'].width=50;sheets['Sources'].column_dimensions['D'].width=53;sheets['Sources'].column_dimensions['E'].width=15;sheets['Sources'].column_dimensions['F'].width=48
    for r in range(4,sheets['Sources'].max_row+1):sheets['Sources'].row_dimensions[r].height=66
    sheets['Assumptions'].column_dimensions['A'].width=49;sheets['Assumptions'].column_dimensions['G'].width=20;sheets['Assumptions'].row_dimensions[14].height=60
    # Keep the overview and source register together; inputs have two clear pages.
    for n in ['Summary','Financing','Checks','Sources']:sheets[n].page_setup.fitToHeight=1
    sheets['Assumptions'].page_setup.orientation='portrait'
    sheets['Assumptions'].row_breaks.append(Break(id=45))
    sheets['Sources'].page_setup.orientation='portrait'
    path.parent.mkdir(parents=True,exist_ok=True);wb.save(path)
    path.with_suffix('.design-expected.json').write_text(json.dumps(expected,ensure_ascii=False,indent=2),encoding='utf-8')
    return expected

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--inputs',type=Path,required=True);parser.add_argument('--output',type=Path,required=True);parser.add_argument('--spec',type=Path,required=True);args=parser.parse_args()
    if not any(x.name=='try' for x in args.output.resolve().parents):raise ValueError('Candidate must be in project try')
    data=json.loads(args.inputs.read_text(encoding='utf-8'));expected=build(data,args.output,args.spec)
    print(json.dumps([{k:v for k,v in e.items() if k not in ['months','totals']} for e in expected],ensure_ascii=False,indent=2))
