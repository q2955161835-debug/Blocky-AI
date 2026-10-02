import openpyxl,json,hashlib
from pathlib import Path
p=Path('doc/交付/2026-10-03-答辩与商业计划书/Blocky-Studio-财务测算-2026-10-03.xlsx')
w=openpyxl.load_workbook(p,data_only=True);f=openpyxl.load_workbook(p,data_only=False);checks=[];out={}
def ck(sh,r,c,v):
 actual=w[sh].cell(r,c).value;delta=abs(actual-v) if isinstance(actual,(int,float)) else None
 checks.append({'cell':f'{sh}!{openpyxl.utils.get_column_letter(c)}{r}','expected':v,'actual':actual,'pass':(delta is not None and delta<.000001) or actual==v})
for si,name in enumerate(['基准','悲观','积极']):
 a={r:w['Assumptions'].cell(r,5+si).value for r in range(10,33)};rr=18*si;cr=24*si
 weights=[w['Revenue'].cell(7+rr,c).value for c in range(4,40)]
 ns=[weights[m]/sum(weights[m//12*12:m//12*12+12])*a[10+m//12] for m in range(36)]
 rev=[n*a[13] for n in ns];ar=[v*(1-a[28]) for v in rev];cash=a[26];months=[]
 for m,n in enumerate(ns):
  hours=n*a[14];internal=min(hours,max(0,a[30]-a[31]));external=max(0,hours-internal);receipt=rev[m]*a[28]+(ar[m-1] if m else 0)
  model=n*(a[17]*a[19]+a[18]*a[20]);pay=internal*a[15];ext=external*a[16];other=n*a[21];direct=model+pay+ext+other;tax=rev[m]*a[29];cap=a[24] if m==0 else 0;dep=a[24]/a[25] if m<a[25] else 0;total=direct+a[22]+a[23]+tax+dep;expense=total-dep;net=rev[m]-total;flow=receipt-expense-cap;cash+=flow
  vals={'n':n,'rev':rev[m],'ar':ar[m],'hours':hours,'internal':internal,'external':external,'receipt':receipt,'model':model,'pay':pay,'ext':ext,'other':other,'direct':direct,'tax':tax,'cap':cap,'dep':dep,'total':total,'expense':expense,'net':net,'flow':flow,'cash':cash};months.append(vals)
  for r,v in {8:n,9:rev[m],10:receipt,11:ar[m],12:hours,13:internal,14:external,15:max(0,a[30]-a[31])+a[32]-hours,16:abs(n-round(n))}.items():ck('Revenue',r+rr,m+4,v)
  for r,v in {7:model,8:pay,9:ext,10:other,11:direct,12:a[22],13:a[23],14:tax,15:cap,16:dep,17:total,18:expense,19:net,20:a[31],21:max(0,a[30]-a[31]),22:a[32]}.items():ck('Costs',r+cr,m+4,v)
 low=min(x['cash'] for x in months);fund=max(0,a[27]-low);peak=max(0,-low)
 for m,x in enumerate(months):
  for r,v in {7:x['rev'],8:x['receipt'],9:x['expense'],10:x['cap'],11:x['flow'],12:x['cash'],13:fund if m==0 else 0,14:x['cash']+fund,15:a[27],16:x['cash']+fund-a[27]}.items():ck('Cash Flow',r+cr,m+4,v)
 for r,v in {6:a[26],7:peak,8:fund,9:a[27],10:low,11:a[26]+fund}.items():ck('Financing',r,4+si,v)
 years=[]
 for y in range(3):
  ms=months[y*12:y*12+12];s=lambda k:sum(x[k] for x in ms);c=4+si*3+y;profit=s('net');cum=sum(x['net'] for x in months[:(y+1)*12]);dep=sum(x['dep'] for x in months[:(y+1)*12]);end=ms[-1]['cash'];ars=ms[-1]['ar'];assets=end+fund+ars+a[24]-dep;equity=a[26]+fund+cum
  for r,v in {8:s('rev'),9:s('direct'),10:s('rev')-s('direct'),11:a[22]*12,12:a[23]*12,13:s('tax'),14:s('dep'),15:s('total'),16:profit}.items():ck('Income Statement',r,c,v)
  for r,v in {8:end+fund,9:ars,10:a[24],11:dep,12:a[24]-dep,13:assets,15:0,16:a[26],17:fund,18:cum,19:equity,21:assets-equity}.items():ck('Balance Sheet',r,c,v)
  for r,v in {8:s('rev'),9:s('total'),10:profit,11:end,12:end+fund,13:s('n'),14:s('hours'),15:s('external'),16:s('direct'),17:s('rev')-s('direct'),18:s('tax')}.items():ck('Summary',r,c,v)
  years.append({'year':2027+y,'revenue':s('rev'),'cost':s('total'),'profit':profit,'cash':end,'external_hours':s('external')})

 for r,v in {6:0,7:0,8:0,9:0,10:min(max(0,a[30]-a[31])+a[32]-x['hours'] for x in months),11:a[30]-a[31],12:0,13:'通过',14:'通过',15:0}.items():ck('Checks',r,4+si,v)
 out[name]={'years':years,'peak_deficit':peak,'funding_with_reserve':fund,'peak_month':f'{2027+months.index(min(months,key=lambda x:x["cash"]))//12}-{months.index(min(months,key=lambda x:x["cash"]))%12+1:02}','max_external_month':max(x['external'] for x in months)}
formula=[(s.title,c.coordinate,w[s.title][c.coordinate].value) for s in f for row in s for c in row if c.data_type=='f'];errors=[(s.title,c.coordinate,c.value) for s in w for row in s for c in row if c.data_type=='e'];result={'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'checks_count':len(checks),'failures':[x for x in checks if not x['pass']],'formula_count':len(formula),'missing_cache':[(s,c) for s,c,v in formula if v is None],'excel_errors':errors,'results':out,'checks':checks}
Path('try/bp-independent-review/financial-recompute.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps({k:v for k,v in result.items() if k!='checks'},ensure_ascii=False,indent=2))
